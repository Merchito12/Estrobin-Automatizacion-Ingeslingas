"""Regresión con los archivos reales de julio (se omite si no están en tests/fixtures/).

HELISA 'VENTAS HELISA JULIO 2026.xls'  →  debe reproducir el Excel que la plataforma YA aceptó.
Se compara semánticamente (hojas, encabezados, valores, tipos, formatos, estilos, columnas ocultas),
no byte a byte.
"""
from datetime import timedelta
from decimal import Decimal

import openpyxl
import pytest

from app.services.file_reader import read_source
from app.services.pipeline import process_upload
from tests.conftest import GOLDEN_FIXTURE, HELISA_FIXTURE

pytestmark = pytest.mark.skipif(
    not (HELISA_FIXTURE.exists() and GOLDEN_FIXTURE.exists()),
    reason="Faltan los archivos reales en tests/fixtures/ (son privados y no se versionan)",
)

EXPECTED_TOTAL = Decimal("476641448.08")


def _border(cell):
    border = cell.border
    return tuple(getattr(getattr(border, side), "style", None) for side in ("left", "right", "top", "bottom"))


@pytest.fixture
def july(store, profile):
    pytest.importorskip("xlrd")
    result = process_upload(HELISA_FIXTURE.read_bytes(), ".xls", None, profile, store)
    assert result.ok, result.payload["errors"]
    job, filename = store.get(result.payload["jobId"])
    return result.payload, openpyxl.load_workbook(job.output_path), filename


def test_xls_is_read_with_xlrd():
    pytest.importorskip("xlrd")
    source = read_source(HELISA_FIXTURE.read_bytes(), ".xls")
    assert source.fmt == "xls" and source.sheets[0].rows[0][2] == "Doc. Fuente"


def test_july_summary(july):
    payload, _, filename = july
    assert payload["recordsFound"] == payload["validRecords"] == 57
    assert Decimal(payload["summary"]["totalValueText"]) == EXPECTED_TOTAL
    assert payload["summary"]["totalValue"] == 476641448.08
    assert filename == payload["outputFilename"] == "PlantillaDeclaracionVentas JULIO 2026.xlsx"
    assert payload["summary"]["minDate"] == "01/07/2026" and payload["warnings"] == []
    assert payload["stats"]["skippedTotalRows"] == 2


def test_july_field_rules(july):
    _, workbook, _ = july
    ws = workbook["Declaración de ventas"]
    rows = [r for r in ws.iter_rows(min_row=2, max_row=ws.max_row) if r[5].value is not None]
    assert len(rows) == 57
    for r in rows:
        a, b, c, d, e, f, g, h, i, j, k, l = r[:12]
        assert b.value == "nitco" and e.value == "co"
        assert isinstance(c.value, str) and c.value.isdigit()
        assert g.data_type == "d" and h.data_type == "d"
        assert h.value - g.value == timedelta(days=90)
        assert i.value == 90 and isinstance(i.value, int)
        assert isinstance(j.value, (int, float))
        assert a.value is None and k.value is None and l.value is None
    assert sum(Decimal(repr(r[9].value)) for r in rows) == EXPECTED_TOTAL


def test_july_matches_accepted_golden_file(july):
    _, workbook, _ = july
    golden = openpyxl.load_workbook(GOLDEN_FIXTURE)
    assert workbook.sheetnames == golden.sheetnames

    ws, gs = workbook["Declaración de ventas"], golden["Declaración de ventas"]
    assert [c.value for c in ws[1]][:12] == [c.value for c in gs[1]][:12]          # encabezados exactos

    golden_rows = [r for r in range(2, gs.max_row + 1) if gs.cell(r, 6).value is not None]
    assert len(golden_rows) == 57
    for r in golden_rows:
        for col in range(1, 13):
            mine, theirs = ws.cell(r, col), gs.cell(r, col)
            assert mine.value == theirs.value, f"{mine.coordinate}: {mine.value!r} != {theirs.value!r}"
            if theirs.value is not None:
                assert type(mine.value) is type(theirs.value), mine.coordinate
                assert mine.data_type == theirs.data_type, mine.coordinate

    # Estilo: cada fila de datos debe parecerse a la fila prototipo (fila 2) del archivo aceptado.
    for r in golden_rows:
        for col in range(1, 13):
            mine, proto = ws.cell(r, col), gs.cell(2, col)
            assert mine.number_format == proto.number_format, mine.coordinate
            assert mine.alignment.horizontal == proto.alignment.horizontal, mine.coordinate
            assert mine.alignment.wrap_text == proto.alignment.wrap_text, mine.coordinate
            assert _border(mine) == _border(proto), mine.coordinate

    for name in golden.sheetnames:
        if name != "Declaración de ventas":
            assert list(workbook[name].iter_rows(values_only=True)) == list(golden[name].iter_rows(values_only=True))
    for letter, dim in gs.column_dimensions.items():
        assert bool(ws.column_dimensions[letter].hidden) == bool(dim.hidden), letter
    assert ws.column_dimensions["D"].width == pytest.approx(gs.column_dimensions["D"].width)
    assert ws.row_dimensions[1].height == gs.row_dimensions[1].height
    assert [c.font.b for c in ws[1][:12]] == [c.font.b for c in gs[1][:12]]
    assert [c.fill.fgColor.rgb for c in ws[1][:12]] == [c.fill.fgColor.rgb for c in gs[1][:12]]
