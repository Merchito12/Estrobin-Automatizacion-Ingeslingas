import dataclasses
import zipfile
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest

from app.models import OK, Record
from app.services import excel_writer, validator
from app.services.postprocess import use_shared_strings
from app.utils.nit import normalize_nit


def make_record(i: int, *, name="CLIENTE S.A.S.", invoice=None, amount="1000.5", issue=date(2026, 8, 3),
                nit="900123456") -> Record:
    rec = Record(source_row=i + 2)
    rec.nit = normalize_nit(nit)
    rec.name = name
    rec.invoice = invoice or f"FECV{17000 + i}"
    rec.issue_date = issue
    rec.due_date = date.fromordinal(issue.toordinal() + 90)
    rec.amount = Decimal(amount)
    rec.term, rec.document_type, rec.country = 90, "nitco", "co"
    rec.status = {k: OK for k in ("nit", "name", "invoice", "issue_date", "amount")}
    return rec


def build(tmp_path, profile, records, shared=True):
    path = tmp_path / "out.xlsx"
    excel_writer.build_workbook(records, profile, path)
    if shared:
        use_shared_strings(path)
    total = sum((r.amount for r in records), Decimal(0))
    return path, total


def test_template_headers_match_profile_exactly(profile):
    workbook = excel_writer.load_template(profile)
    ws = workbook[profile.sheet]
    assert [ws.cell(1, c).value for c in range(1, 13)] == profile.headers
    assert ws["A1"].value == "COD. \nSOLUNION"        # espacios y saltos de línea intactos
    assert ws["B1"].value == "TIPO DOCUMENTO "
    assert ws["C1"].value == "CÓDIGO DOCUMENTO "
    assert workbook.sheetnames == ["Página de ayuda", "Declaración de ventas", "País", "Tipo documento", "Plazos"]


def test_template_has_no_sales_data(profile):
    ws = excel_writer.load_template(profile)[profile.sheet]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        assert all(cell.value is None for cell in row)


def test_generated_workbook_passes_full_verification(tmp_path, profile):
    records = [make_record(i, amount=str(1000 + i) + ".25") for i in range(10)]
    path, total = build(tmp_path, profile, records)
    assert validator.verify_workbook(path, records, profile, total) == []


def test_final_cell_types(tmp_path, profile):
    records = [make_record(0, amount="14208600"), make_record(1, amount="761926.06")]
    path, _ = build(tmp_path, profile, records)
    ws = openpyxl.load_workbook(path)[profile.sheet]
    row = {c.column_letter: c for c in ws[2]}

    assert (row["B"].value, row["B"].data_type) == ("nitco", "s")
    assert (row["C"].value, row["C"].data_type) == ("9001234561", "s") or row["C"].data_type == "s"
    assert isinstance(row["C"].value, str)
    assert (row["E"].value, row["E"].data_type) == ("co", "s")
    assert isinstance(row["F"].value, str) and row["F"].value.startswith("FECV")
    for letter in ("G", "H"):
        assert isinstance(row[letter].value, datetime) and row[letter].data_type == "d"
    assert row["G"].number_format == "m/d/yy"
    assert row["H"].number_format == "d/mm/yyyy;@"
    assert (row["H"].value - row["G"].value).days == 90
    assert row["I"].value == 90 and isinstance(row["I"].value, int) and row["I"].data_type == "n"
    assert row["J"].value == 14208600 and isinstance(row["J"].value, int)
    assert ws["J3"].value == 761926.06 and isinstance(ws["J3"].value, float)
    for letter in ("A", "K", "L"):
        assert row[letter].value is None


def test_old_template_data_never_remains_when_fewer_records(tmp_path, profile):
    """El workbook tiene 85 filas con estilo; con 3 registros las demás deben quedar vacías."""
    records = [make_record(i) for i in range(3)]
    path, total = build(tmp_path, profile, records)
    ws = openpyxl.load_workbook(path)[profile.sheet]
    assert ws.max_row >= 86
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        assert all(c.value is None for c in row)
    assert validator.verify_workbook(path, records, profile, total) == []


def test_more_records_than_prestyled_rows_copy_prototype_style(tmp_path, profile):
    records = [make_record(i, amount="10.5") for i in range(120)]
    path, total = build(tmp_path, profile, records)
    ws = openpyxl.load_workbook(path)[profile.sheet]
    assert ws.max_row == 121
    for r in (2, 86, 87, 121):
        assert ws.cell(r, 7).number_format == "m/d/yy"
        assert ws.cell(r, 8).number_format == "d/mm/yyyy;@"
        assert ws.cell(r, 7).data_type == "d"
    assert validator.verify_workbook(path, records, profile, total) == []


def test_formula_like_text_is_stored_as_text_not_formula(tmp_path, profile):
    records = [make_record(0, name="=HYPERLINK(\"http://x\",\"a\")", invoice="=1+1")]
    path, total = build(tmp_path, profile, records)
    ws = openpyxl.load_workbook(path)[profile.sheet]
    assert ws["D2"].data_type == "s" and ws["D2"].value.startswith("=")
    assert ws["F2"].data_type == "s" and ws["F2"].value == "=1+1"
    assert validator.verify_workbook(path, records, profile, total) == []


def test_auxiliary_sheets_and_hidden_columns_preserved(tmp_path, profile):
    path, _ = build(tmp_path, profile, [make_record(0)])
    out = openpyxl.load_workbook(path)
    template = excel_writer.load_template(profile)
    assert out.sheetnames == template.sheetnames
    for name in profile.auxiliary_sheets:
        assert list(out[name].iter_rows(values_only=True)) == list(template[name].iter_rows(values_only=True))
    dims = out[profile.sheet].column_dimensions
    assert dims["N"].hidden and dims["T"].hidden


def test_shared_strings_conversion_keeps_every_value(tmp_path, profile):
    records = [make_record(i, name=f"  CLIENTE {i}  ".strip()) for i in range(5)]
    raw_path, _ = build(tmp_path, profile, records, shared=False)
    before = [list(ws.iter_rows(values_only=True)) for ws in openpyxl.load_workbook(raw_path).worksheets]
    with zipfile.ZipFile(raw_path) as z:
        assert "xl/sharedStrings.xml" not in z.namelist()
    assert use_shared_strings(raw_path) is True
    with zipfile.ZipFile(raw_path) as z:
        assert "xl/sharedStrings.xml" in z.namelist()
        assert b"inlineStr" not in b"".join(z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/"))
    after = [list(ws.iter_rows(values_only=True)) for ws in openpyxl.load_workbook(raw_path).worksheets]
    assert before == after
    assert use_shared_strings(raw_path) is False       # idempotente


def test_header_with_line_break_survives_shared_strings(tmp_path, profile):
    path, _ = build(tmp_path, profile, [make_record(0)])
    ws = openpyxl.load_workbook(path)[profile.sheet]
    assert ws["A1"].value == "COD. \nSOLUNION"


@pytest.mark.parametrize("tamper, expected", [
    (lambda ws: ws.__setitem__("G2", "01/08/2026"), "FECHA EMISIÓN"),
    (lambda ws: ws.__setitem__("I2", "90"), "PLAZO"),
    (lambda ws: ws.__setitem__("J2", "1.000,50"), "VALOR"),
    (lambda ws: ws.__setitem__("C2", 9001234561), "CÓDIGO DOCUMENTO"),
    (lambda ws: ws.__setitem__("A40", "residuo"), "residual"),
    (lambda ws: ws.__setitem__("H2", "=G2+90"), "FECHA VENCIMIENTO"),
])
def test_verify_detects_tampered_workbooks(tmp_path, profile, tamper, expected):
    records = [make_record(0)]
    path, total = build(tmp_path, profile, records)
    workbook = openpyxl.load_workbook(path)
    tamper(workbook[profile.sheet])
    workbook.save(path)
    problems = validator.verify_workbook(path, records, profile, total)
    assert problems and any(expected in p for p in problems), problems


def test_verify_detects_wrong_total(tmp_path, profile):
    records = [make_record(0, amount="100")]
    path, _ = build(tmp_path, profile, records)
    assert any("suma" in p for p in validator.verify_workbook(path, records, profile, Decimal("101")))


def test_template_error_when_header_differs(tmp_path, profile):
    bad = tmp_path / "bad.xlsx"
    workbook = excel_writer.load_template(profile)
    workbook[profile.sheet]["D1"] = "OTRO ENCABEZADO"
    workbook.save(bad)
    with pytest.raises(excel_writer.TemplateError):
        excel_writer.load_template(dataclasses.replace(profile, template_path=bad))


def test_template_error_when_template_is_missing(tmp_path, profile):
    with pytest.raises(excel_writer.TemplateError):
        excel_writer.load_template(dataclasses.replace(profile, template_path=tmp_path / "no-existe.xlsx"))
