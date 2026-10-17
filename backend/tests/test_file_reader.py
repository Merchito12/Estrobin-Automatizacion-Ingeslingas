import io
import zipfile
from datetime import date, datetime

import openpyxl
import pytest

from app.services.file_reader import FileReadError, decode_text, detect_delimiter, read_source
from tests.conftest import make_xlsx


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_csv_delimiters_are_detected(delimiter):
    text = delimiter.join(["NIT", "Cliente", "Valor"]) + "\n" + delimiter.join(["900123456", "ACME", "10"]) + "\n"
    source = read_source(text.encode(), ".csv")
    assert source.fmt == "csv"
    assert source.sheets[0].rows == [["NIT", "Cliente", "Valor"], ["900123456", "ACME", "10"]]


def test_delimiter_inside_quotes_is_ignored():
    text = 'NIT;Cliente;Valor\n900123456;"ACME, S.A.";"1.234,56"\n'
    assert detect_delimiter(text) == ";"
    assert read_source(text.encode(), ".csv").sheets[0].rows[1] == ["900123456", "ACME, S.A.", "1.234,56"]


def test_utf8_bom_and_latin1_are_decoded():
    assert decode_text("﻿NIT;Ñandú".encode("utf-8")) == "NIT;Ñandú"
    assert decode_text("Construcción".encode("latin-1")) == "Construcción"
    assert decode_text("Construcción".encode("utf-8")) == "Construcción"


def test_csv_empty_cells_become_none_and_trailing_columns_are_trimmed():
    rows = read_source(b"a,b,c\n1,,\n,,\n", ".csv").sheets[0].rows
    assert rows == [["a", "b", "c"], ["1"], []]


def test_xlsx_reads_values_dates_and_all_sheets():
    data = make_xlsx([["Fecha", "Valor"], [datetime(2026, 7, 1), 12.5]], "Datos", {"Otra": [["x"]]})
    source = read_source(data, ".xlsx")
    assert source.fmt == "xlsx" and [s.name for s in source.sheets] == ["Datos", "Otra"]
    assert source.sheets[0].rows[1] == [datetime(2026, 7, 1), 12.5]


def test_formulas_are_never_executed():
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws["A1"], ws["A2"], ws["A3"] = "Valor", "=1+1", '=HYPERLINK("http://evil","x")'
    buffer = io.BytesIO()
    workbook.save(buffer)
    rows = read_source(buffer.getvalue(), ".xlsx").sheets[0].rows
    assert rows[0] == ["Valor"] and rows[1:] == [[], []]       # sin valor en caché: no se evalúa nada


def test_excel_error_values_become_none():
    data = make_xlsx([["a", "b"], ["#N/A", "ok"]])
    assert read_source(data, ".xlsx").sheets[0].rows[1] == [None, "ok"]


def test_html_table_is_supported():
    html = b"<html><body><table><tr><td>NIT</td><td>Valor</td></tr><tr><td>900</td><td>1&nbsp;000</td></tr></table></body></html>"
    source = read_source(html, ".xls")
    assert source.fmt == "html" and source.sheets[0].rows[0] == ["NIT", "Valor"]


@pytest.mark.parametrize("data", [b"", b"PK\x03\x04basura", b"%PDF-1.7 \x00\x01\x02", bytes(range(256)) * 4])
def test_unreadable_content_raises_clear_error(data):
    with pytest.raises(FileReadError) as info:
        read_source(data, ".xlsx")
    assert info.value.message


def test_zip_without_workbook_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "hola")
    with pytest.raises(FileReadError):
        read_source(buffer.getvalue(), ".xlsx")


def test_real_date_cells_are_returned_as_dates():
    rows = read_source(make_xlsx([["f"], [date(2026, 7, 1)]]), ".xlsx").sheets[0].rows
    assert isinstance(rows[1][0], (datetime, date))
