"""Flujo completo (lectura → mapeo → validación → Excel → verificación) con datos sintéticos."""
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from app.services.pipeline import process_upload
from tests.conftest import HELISA_HEADER, helisa_rows, make_xlsx, simple_rows


def run(data, store, profile, ext=".xlsx", portal_error=None):
    return process_upload(data, ext, portal_error, profile, store)


def test_helisa_like_xlsx_end_to_end(store, profile):
    result = run(make_xlsx(helisa_rows(6), "Consulta"), store, profile)
    assert result.ok and result.http_status == 200
    payload = result.payload
    assert payload["recordsFound"] == payload["validRecords"] == 6
    assert payload["outputFilename"] == "PlantillaDeclaracionVentas JULIO 2026.xlsx"
    assert payload["errors"] == []
    assert payload["summary"]["period"] == "JULIO 2026"
    assert payload["summary"]["minDate"] == "01/07/2026" and payload["summary"]["maxDate"] == "06/07/2026"
    assert len(payload["preview"]) == 6
    targets = {m["target"]: m for m in payload["mapping"]}
    assert targets["VALOR DE LA FACTURA"]["source"] == "Vr Neto" and targets["VALOR DE LA FACTURA"]["confidence"] == 1.0
    assert targets["TIPO DOCUMENTO"]["value"] == "nitco"

    found = store.get(payload["jobId"])
    assert found is not None
    ws = openpyxl.load_workbook(found[0].output_path)["Declaración de ventas"]
    assert ws.cell(2, 6).value == "FECV17000"
    assert Decimal(str(payload["summary"]["totalValueText"])) == sum(
        (Decimal(repr(round(1000000 + i * 12345.67, 2))) for i in range(6)), Decimal(0))


def test_total_is_exact_decimal_sum(store, profile):
    rows = simple_rows(dates=[date(2026, 8, 3)] * 3, amounts=["0.1", "0.2", "0.3"],
                       invoices=["A1", "A2", "A3"])
    payload = run(make_xlsx(rows), store, profile).payload
    assert payload["ok"] and payload["summary"]["totalValueText"] == "0.6"


def test_preview_is_limited_to_15_rows_but_all_records_are_written(store, profile):
    rows = simple_rows(dates=[date(2026, 8, 3)] * 40, amounts=[10] * 40, invoices=[f"F{i}" for i in range(40)])
    payload = run(make_xlsx(rows), store, profile).payload
    assert payload["recordsFound"] == 40 and len(payload["preview"]) == 15
    ws = openpyxl.load_workbook(store.get(payload["jobId"])[0].output_path)["Declaración de ventas"]
    assert ws.cell(41, 6).value == "F39" and ws.cell(42, 6).value is None


def test_rejected_platform_file_is_only_a_data_source(store, profile):
    """Un archivo con la estructura de la plantilla pero con errores (fechas texto, vencimiento mal) se reconstruye."""
    rows = [profile.headers,
            [None, "nitco", "9012339571", "E Y R CIMENTACIONES S.A.S.", "co", "FECV17451", "01/06/2026",
             "=G2+90", "90", "1.789.760", None, None],
            [None, "nitco", "800166199", "DELTEC S.A.", "co", "FECV17452", 46174, "30/08/2026", 90, "1115030", None, None]]
    result = run(make_xlsx(rows, "Declaración de ventas", {"País": [["Valor a introducir", "Descripción"], ["co", "Colombia"]]}),
                 store, profile, portal_error="La fila 2 de la columna G está incorrecta, debe quedar en formato de fecha corta")
    assert result.ok, result.payload["errors"]
    payload = result.payload
    assert payload["outputFilename"] == "PlantillaDeclaracionVentas JUNIO 2026.xlsx"
    assert any(h["code"] == "dates" for h in payload["portalErrorHints"])
    ws = openpyxl.load_workbook(store.get(payload["jobId"])[0].output_path)["Declaración de ventas"]
    assert ws["C2"].value == "9012339571" and ws["C3"].value == "8001661991"
    assert ws["G2"].value.date() == date(2026, 6, 1) and ws["H2"].value.date() == date(2026, 8, 30)
    assert ws["J2"].value == 1789760 and ws["I3"].value == 90


def test_errors_block_generation_and_leave_no_files(store, profile):
    rows = simple_rows(dates=[date(2026, 8, 3)] * 2, invoices=["DUP", "DUP"])
    result = run(make_xlsx(rows), store, profile)
    assert not result.ok
    payload = result.payload
    assert payload["errorCount"] == 1 and "duplicada" in payload["errors"][0]["message"]
    assert payload["errors"][0]["row"] == 3
    assert "jobId" not in payload
    assert list(store.root.iterdir()) == []


def test_missing_required_column_message(store, profile):
    rows = [["NIT", "Cliente", "Número factura", "Fecha factura"], ["900123456", "A", "F1", "01/08/2026"]]
    payload = run(make_xlsx(rows), store, profile).payload
    assert not payload["ok"]
    assert payload["errors"][0]["message"] == "No fue posible identificar automáticamente la columna VALOR DE LA FACTURA."
    assert "detectedHeaders" in payload


def test_file_without_headers_is_rejected_clearly(store, profile):
    payload = run(make_xlsx([["a", "b"], [1, 2]]), store, profile).payload
    assert not payload["ok"] and "encabezados" in payload["errors"][0]["message"]


def test_csv_semicolon_with_colombian_amounts(store, profile):
    text = "NIT;Cliente;Número factura;Fecha factura;Valor Neto\r\n" \
           "900.123.456;ACME S.A.S.;F1;01/08/2026;\"1.234.567,89\"\r\n" \
           "800222333;Beta Ltda;F2;02/08/2026;$ 2.000.000,10\r\n"
    result = run(text.encode("utf-8-sig"), store, profile, ext=".csv")
    assert result.ok, result.payload["errors"]
    assert result.payload["summary"]["totalValueText"] == "3234567.99"


def test_csv_latin1_comma_with_accents(store, profile):
    text = "NIT,Cliente,Factura,Fecha,Vr Neto\n900123456,Ñandú Construcción S.A.,F1,2026-08-01,1000\n"
    result = run(text.encode("latin-1"), store, profile, ext=".csv")
    assert result.ok
    ws = openpyxl.load_workbook(store.get(result.payload["jobId"])[0].output_path)["Declaración de ventas"]
    assert ws["D2"].value == "Ñandú Construcción S.A."


def test_garbage_file_is_a_clear_error_not_a_crash(store, profile):
    result = run(b"PK\x03\x04 esto no es un excel", store, profile)
    assert not result.ok and result.http_status == 422
    assert "dañado" in result.payload["errors"][0]["message"] or "válido" in result.payload["errors"][0]["message"]
    result = run(b"", store, profile)
    assert not result.ok


def test_html_exported_as_xls_is_read(store, profile):
    html = ("<html><body><table><tr><th>NIT</th><th>Cliente</th><th>Factura</th><th>Fecha</th><th>Vr Neto</th></tr>"
            "<tr><td>900123456</td><td>ACME</td><td>F1</td><td>01/08/2026</td><td>1.500,50</td></tr></table></body></html>")
    result = run(html.encode("utf-8"), store, profile, ext=".xls")
    assert result.ok and result.payload["summary"]["totalValueText"] == "1500.50"


def test_nit_with_wrong_dv_produces_warning_but_succeeds(store, profile):
    rows = [["CÓDIGO DOCUMENTO", "Cliente", "Factura", "Fecha", "Vr Neto"],
            ["9012339570", "ACME", "F1", "01/08/2026", 10]]
    payload = run(make_xlsx(rows), store, profile).payload
    assert payload["ok"] and any(w["code"] == "nit_dv_corrected" for w in payload["warnings"])
    ws = openpyxl.load_workbook(store.get(payload["jobId"])[0].output_path)["Declaración de ventas"]
    assert ws["C2"].value == "9012339571"


def test_multiple_periods_warn_and_use_dominant(store, profile):
    rows = simple_rows(dates=[date(2026, 8, 1), date(2026, 8, 2), date(2026, 9, 1)], invoices=["A", "B", "C"])
    payload = run(make_xlsx(rows), store, profile).payload
    assert payload["ok"] and payload["outputFilename"].endswith("AGOSTO 2026.xlsx")
    assert any(w["code"] == "multiple_periods" for w in payload["warnings"])


def test_errors_are_capped_in_the_response(store, profile):
    rows = simple_rows(dates=[date(2026, 8, 1)] * 300, nits=[""] * 300, invoices=[f"F{i}" for i in range(300)])
    payload = run(make_xlsx(rows), store, profile).payload
    assert payload["errorCount"] == 300 and len(payload["errors"]) == 100


def test_helisa_header_constant_is_what_the_fixture_expects():
    assert HELISA_HEADER[2] == "Doc. Fuente"
