"""Lectura de archivos de entrada (.xls, .xlsx, .csv) como tablas de valores crudos.

El contenido nunca se ejecuta: no hay macros, las fórmulas no se evalúan (se leen
sus valores en caché) y los vínculos externos no se siguen.
"""
from __future__ import annotations

import csv
import io
import logging
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from html import unescape
from html.parser import HTMLParser
from typing import Any

logger = logging.getLogger("ingeslingas.reader")

MAX_ROWS = 200_000
MAX_COLUMNS = 200
MAX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024
OLE2_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_SIGNATURE = b"PK\x03\x04"
_EXCEL_ERRORS = {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}
CSV_DELIMITERS = (",", ";", "\t", "|")


class FileReadError(Exception):
    """El archivo no se pudo interpretar. ``message`` es seguro de mostrar al usuario."""

    def __init__(self, message: str, code: str = "unreadable") -> None:
        super().__init__(message)
        self.message = message
        self.code = code


@dataclass
class SourceSheet:
    name: str
    rows: list[list[Any]]


@dataclass
class SourceFile:
    fmt: str                        # xls | xlsx | csv | html
    sheets: list[SourceSheet] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Punto de entrada
# --------------------------------------------------------------------------- #
def read_source(data: bytes, extension: str) -> SourceFile:
    """Lee ``data`` según su contenido real (la extensión solo orienta)."""
    if not data:
        raise FileReadError("El archivo está vacío.", "empty_file")

    if data.startswith(OLE2_SIGNATURE):
        return _read_xls(data)
    if data.startswith(ZIP_SIGNATURE):
        return _read_xlsx(data)

    head = data[:4096].lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if head.startswith((b"<html", b"<!doctype", b"<table", b"<?xml", b"<meta", b"<head", b"<body")):
        return _read_html(data)
    if extension == ".csv" or _looks_like_text(data):
        return _read_csv(data)
    raise FileReadError(
        "No se pudo reconocer el contenido del archivo como Excel (.xls/.xlsx) ni como CSV.",
        "unknown_format",
    )


def _looks_like_text(data: bytes) -> bool:
    sample = data[:8192]
    if b"\x00" in sample:
        return False
    return True


# --------------------------------------------------------------------------- #
# .xlsx
# --------------------------------------------------------------------------- #
def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
            names = {info.filename for info in archive.infolist()}
    except zipfile.BadZipFile as exc:
        raise FileReadError("El archivo .xlsx está dañado o no es un libro de Excel válido.") from exc
    if total > MAX_UNCOMPRESSED_BYTES:
        raise FileReadError("El archivo descomprimido es demasiado grande.", "too_large")
    if "xl/workbook.xml" not in names:
        raise FileReadError("El archivo comprimido no es un libro de Excel (.xlsx) válido.")


def _clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        stripped = value.strip()
        if stripped in _EXCEL_ERRORS:
            return None
        return value if stripped else None
    return value


def _trim_row(values: list[Any]) -> list[Any]:
    end = len(values)
    while end and values[end - 1] is None:
        end -= 1
    return values[: min(end, MAX_COLUMNS)]


def _read_xlsx(data: bytes) -> SourceFile:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    _check_zip(data)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True, keep_links=False
        )
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise FileReadError("No se pudo abrir el libro de Excel (.xlsx).") from exc
    except Exception as exc:  # openpyxl lanza tipos variados con archivos corruptos
        logger.warning("openpyxl no pudo abrir el archivo: %s", type(exc).__name__)
        raise FileReadError("No se pudo abrir el libro de Excel (.xlsx).") from exc

    sheets: list[SourceSheet] = []
    try:
        for worksheet in workbook.worksheets:
            try:
                worksheet.reset_dimensions()
            except AttributeError:  # pragma: no cover - versiones antiguas
                pass
            rows: list[list[Any]] = []
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS:
                    raise FileReadError(f"El archivo supera el máximo de {MAX_ROWS:,} filas.", "too_many_rows")
                rows.append(_trim_row([_clean_cell(v) for v in row]))
            sheets.append(SourceSheet(worksheet.title, rows))
    finally:
        workbook.close()
    return SourceFile("xlsx", sheets)


# --------------------------------------------------------------------------- #
# .xls (BIFF)
# --------------------------------------------------------------------------- #
def _read_xls(data: bytes) -> SourceFile:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise FileReadError("El servidor no tiene instalada la librería para leer .xls.", "server") from exc

    try:
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True, ragged_rows=True)
    except xlrd.XLRDError as exc:
        raise FileReadError(
            "No se pudo abrir el archivo .xls (puede estar protegido con contraseña o dañado)."
        ) from exc
    except Exception as exc:
        logger.warning("xlrd no pudo abrir el archivo: %s", type(exc).__name__)
        raise FileReadError("No se pudo abrir el archivo .xls.") from exc

    sheets: list[SourceSheet] = []
    try:
        for index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(index)
            if sheet.nrows > MAX_ROWS:
                raise FileReadError(f"El archivo supera el máximo de {MAX_ROWS:,} filas.", "too_many_rows")
            rows = [
                _trim_row([_xls_value(cell, workbook.datemode) for cell in sheet.row(r)])
                for r in range(sheet.nrows)
            ]
            sheets.append(SourceSheet(sheet.name, rows))
            workbook.unload_sheet(index)
    finally:
        workbook.release_resources()
    return SourceFile("xls", sheets)


def _xls_value(cell: Any, datemode: int) -> Any:
    import xlrd

    ctype, value = cell.ctype, cell.value
    if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
        return None
    if ctype == xlrd.XL_CELL_TEXT:
        return _clean_cell(value)
    if ctype == xlrd.XL_CELL_NUMBER:
        return int(value) if float(value).is_integer() and abs(value) < 1e15 else value
    if ctype == xlrd.XL_CELL_DATE:
        try:
            converted = xlrd.xldate.xldate_as_datetime(value, datemode)
        except (xlrd.xldate.XLDateError, ValueError, OverflowError):
            return None
        return converted if isinstance(converted, datetime) else None
    if ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(value)
    return None


# --------------------------------------------------------------------------- #
# CSV / texto delimitado
# --------------------------------------------------------------------------- #
def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    try:
        return data.decode("cp1252")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def detect_delimiter(text: str) -> str:
    lines = [ln for ln in text.splitlines()[:50] if ln.strip()]
    if not lines:
        return ","
    best, best_score = ",", -1.0
    for delimiter in CSV_DELIMITERS:
        counts = [_count_outside_quotes(line, delimiter) for line in lines]
        if not counts or max(counts) == 0:
            continue
        # Prefiere el separador más frecuente y consistente entre líneas.
        common = max(set(counts), key=counts.count)
        consistency = counts.count(common) / len(counts)
        score = common * consistency if common > 0 else 0
        if score > best_score:
            best, best_score = delimiter, score
    return best


def _count_outside_quotes(line: str, delimiter: str) -> int:
    count, in_quotes = 0, False
    for ch in line:
        if ch == '"':
            in_quotes = not in_quotes
        elif ch == delimiter and not in_quotes:
            count += 1
    return count


def _read_csv(data: bytes) -> SourceFile:
    text = decode_text(data)
    delimiter = detect_delimiter(text)
    csv.field_size_limit(1_000_000)
    rows: list[list[Any]] = []
    try:
        for index, record in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)):
            if index >= MAX_ROWS:
                raise FileReadError(f"El archivo supera el máximo de {MAX_ROWS:,} filas.", "too_many_rows")
            rows.append(_trim_row([_clean_cell(v) for v in record]))
    except csv.Error as exc:
        raise FileReadError("El archivo CSV tiene un formato inválido.") from exc
    if not any(rows):
        raise FileReadError("El archivo CSV no contiene datos.", "empty_file")
    return SourceFile("csv", [SourceSheet("csv", rows)])


# --------------------------------------------------------------------------- #
# HTML disfrazado de .xls (algunas exportaciones contables)
# --------------------------------------------------------------------------- #
class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._depth = 0

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag == "table":
            self._depth += 1
            if self._depth == 1:
                self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(unescape("".join(self._cell)).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None and self.tables:
            self.tables[-1].append(self._row)
            self._row = None
        elif tag == "table":
            self._depth = max(0, self._depth - 1)

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def _read_html(data: bytes) -> SourceFile:
    parser = _TableParser()
    try:
        parser.feed(decode_text(data))
    except Exception as exc:  # HTMLParser es tolerante, pero por seguridad
        raise FileReadError("No se pudo interpretar el contenido HTML del archivo.") from exc
    tables = [t for t in parser.tables if t]
    if not tables:
        raise FileReadError("El archivo parece HTML y no contiene ninguna tabla.")
    table = max(tables, key=len)
    rows = [_trim_row([cell if cell.strip() else None for cell in row]) for row in table[:MAX_ROWS]]
    return SourceFile("html", [SourceSheet("tabla", rows)])
