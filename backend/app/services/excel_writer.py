"""Construcción del Excel final SIEMPRE sobre la plantilla oficial aceptada.

Nunca se usa ``DataFrame.to_excel``: se abre una copia de la plantilla, se limpian
solo los datos antiguos y se escriben los registros nuevos conservando hojas,
estilos, anchos, columnas ocultas y formatos de fecha.
"""
from __future__ import annotations

from copy import copy
from datetime import datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from ..models import Record
from ..profile import Profile


class TemplateError(Exception):
    """La plantilla base no existe o no coincide con el perfil."""


def load_template(profile: Profile) -> openpyxl.Workbook:
    if not profile.template_path.exists():
        raise TemplateError(f"No se encontró la plantilla base: {profile.template_path.name}")
    workbook = openpyxl.load_workbook(profile.template_path, keep_links=False)
    if profile.sheet not in workbook.sheetnames:
        raise TemplateError(f"La plantilla no contiene la hoja '{profile.sheet}'.")
    for name in profile.auxiliary_sheets:
        if name not in workbook.sheetnames:
            raise TemplateError(f"La plantilla no contiene la hoja auxiliar '{name}'.")
    worksheet = workbook[profile.sheet]
    for column in profile.columns:
        actual = worksheet.cell(profile.header_row, column.index).value
        if actual != column.header:
            raise TemplateError(
                f"El encabezado {column.letter}{profile.header_row} de la plantilla ({actual!r}) "
                f"no coincide con el del perfil ({column.header!r})."
            )
    return workbook


def build_workbook(records: list[Record], profile: Profile, output_path: Path) -> None:
    """Escribe ``records`` en una copia de la plantilla base y la guarda en ``output_path``."""
    workbook = load_template(profile)
    worksheet = workbook[profile.sheet]

    prototype = {
        col: copy(worksheet.cell(profile.style_prototype_row, col)._style)
        for col in range(1, profile.style_columns + 1)
    }
    prototype_height = worksheet.row_dimensions[profile.style_prototype_row].height

    first = profile.first_data_row
    last_needed = first + max(len(records), 1) - 1
    last_row = max(worksheet.max_row, last_needed)

    # 1) Limpia TODA la zona de datos de la plantilla (valores) y uniforma el estilo.
    for row in range(first, last_row + 1):
        for col in range(1, profile.style_columns + 1):
            cell = worksheet.cell(row, col)
            cell.value = None
            cell._style = copy(prototype[col])
        worksheet.row_dimensions[row].height = prototype_height

    # 2) Escribe los registros nuevos con el tipo correcto por celda.
    for offset, record in enumerate(records):
        _write_record(worksheet, first + offset, record, profile)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def _write_record(ws: Worksheet, row: int, rec: Record, profile: Profile) -> None:
    cell = lambda key: ws.cell(row, profile.column(key).index)  # noqa: E731

    _text(cell("solunion_code"), rec.solunion_code)
    _text(cell("document_type"), rec.document_type)
    _text(cell("nit"), rec.nit.full if rec.nit else None)
    _text(cell("name"), rec.name)
    _text(cell("country"), rec.country)
    _text(cell("invoice"), rec.invoice)
    _date(cell("issue_date"), rec.issue_date)
    _date(cell("due_date"), rec.due_date)
    _integer(cell("term"), rec.term)
    _number(cell("amount"), rec.amount)
    _text(cell("debtor_reference"), rec.debtor_reference)
    _text(cell("payment_condition"), rec.payment_condition)


def _text(cell: Any, value: str | None) -> None:
    """Texto real. Nunca una fórmula, aunque el contenido empiece por '='."""
    if value is None or value == "":
        cell.value = None
        return
    cell.value = str(value)
    cell.data_type = "s"


def _date(cell: Any, value: Any) -> None:
    if value is None:
        cell.value = None
        return
    cell.value = datetime.combine(value, time.min)   # fecha real; formato = el de la plantilla


def _integer(cell: Any, value: int | None) -> None:
    cell.value = None if value is None else int(value)


def _number(cell: Any, value: Decimal | None) -> None:
    if value is None:
        cell.value = None
    elif value == value.to_integral_value():
        cell.value = int(value)
    else:
        cell.value = float(value)
