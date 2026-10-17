"""Validaciones en tres niveles: datos, reglas de negocio y workbook final.

* Nivel 1 – ``validate_data``: campos obligatorios, formatos y duplicados.
* Nivel 2 – ``validate_rules``: constantes y derivados cumplen las reglas del perfil.
* Nivel 3 – ``verify_workbook``: se reabre el .xlsx YA GUARDADO y se comprueba
  estructura, tipos y valores celda por celda.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles.numbers import is_date_format

from ..models import EMPTY, ERROR, INVALID, WARNING, Issue, Record
from ..profile import Profile

MAX_RAW_CHARS = 40


def _shown(value: Any) -> str:
    text = str(value)
    return text if len(text) <= MAX_RAW_CHARS else text[: MAX_RAW_CHARS - 1] + "…"


# --------------------------------------------------------------------------- #
# Nivel 1: datos
# --------------------------------------------------------------------------- #
def validate_data(records: list[Record], profile: Profile) -> list[Issue]:
    issues: list[Issue] = []
    first_row_by_invoice: dict[str, int] = {}

    negative_severity = ERROR if profile.rules.get("negative_amount_severity", ERROR) == ERROR else WARNING
    zero_severity = ERROR if profile.rules.get("zero_amount_severity", WARNING) == ERROR else WARNING

    for rec in records:
        row = rec.source_row

        status = rec.status.get("nit")
        if status == EMPTY:
            issues.append(Issue(ERROR, "nit_empty", "NIT/documento vacío.", row, "CÓDIGO DOCUMENTO"))
        elif status == INVALID:
            issues.append(Issue(
                ERROR, "nit_invalid",
                f"NIT/documento inválido ({rec.detail.get('nit', 'formato no reconocido')}): "
                f"'{_shown(rec.raw.get('nit'))}'.", row, "CÓDIGO DOCUMENTO"))

        if rec.status.get("name") == EMPTY:
            issues.append(Issue(ERROR, "name_empty", "Nombre/razón social vacío.", row, "NOMBRE / RAZÓN SOCIAL"))

        if rec.status.get("invoice") == EMPTY:
            issues.append(Issue(ERROR, "invoice_empty", "Número de factura vacío.", row, "NÚMERO FACTURA"))
        elif rec.invoice:
            first = first_row_by_invoice.get(rec.invoice)
            if first is None:
                first_row_by_invoice[rec.invoice] = row
            else:
                issues.append(Issue(
                    ERROR, "invoice_duplicate",
                    f"Factura duplicada: {rec.invoice} (ya aparece en la fila {first}).",
                    row, "NÚMERO FACTURA"))

        status = rec.status.get("issue_date")
        if status == EMPTY:
            issues.append(Issue(ERROR, "date_empty", "Fecha de emisión vacía.", row, "FECHA EMISIÓN"))
        elif status == INVALID:
            issues.append(Issue(
                ERROR, "date_invalid",
                f"Fecha de emisión inválida: '{_shown(rec.raw.get('issue_date'))}'.", row, "FECHA EMISIÓN"))

        status = rec.status.get("amount")
        if status == EMPTY:
            issues.append(Issue(ERROR, "amount_empty", "Valor de la factura vacío.", row, "VALOR DE LA FACTURA"))
        elif status == INVALID:
            issues.append(Issue(
                ERROR, "amount_invalid",
                f"Valor de la factura inválido: '{_shown(rec.raw.get('amount'))}'.", row, "VALOR DE LA FACTURA"))
        elif rec.amount is not None:
            if rec.amount < 0:
                issues.append(Issue(
                    negative_severity, "amount_negative",
                    f"Valor negativo inesperado: {rec.amount}.", row, "VALOR DE LA FACTURA"))
            elif rec.amount == 0:
                issues.append(Issue(zero_severity, "amount_zero", "El valor de la factura es cero.",
                                    row, "VALOR DE LA FACTURA"))
    return issues


# --------------------------------------------------------------------------- #
# Nivel 2: reglas
# --------------------------------------------------------------------------- #
def validate_rules(records: list[Record], profile: Profile) -> list[Issue]:
    issues: list[Issue] = []
    offset = int(profile.rules["due_offset_days"])
    document_type = str(profile.constants["document_type"])
    country = str(profile.constants["country"])
    term = int(profile.constants["term_days"])

    for rec in records:
        row = rec.source_row
        if rec.document_type != document_type:
            issues.append(Issue(ERROR, "rule_document_type", f"TIPO DOCUMENTO debe ser '{document_type}'.", row))
        if rec.country != country:
            issues.append(Issue(ERROR, "rule_country", f"PAÍS debe ser '{country}'.", row))
        if not isinstance(rec.term, int) or isinstance(rec.term, bool) or rec.term != term:
            issues.append(Issue(ERROR, "rule_term", f"PLAZO DE VENTA debe ser el entero {term}.", row))
        if rec.issue_date is not None:
            if rec.due_date is None or (rec.due_date - rec.issue_date).days != offset:
                issues.append(Issue(
                    ERROR, "rule_due_date", f"FECHA VENCIMIENTO debe ser FECHA EMISIÓN + {offset} días.", row))
        if rec.amount is not None and not isinstance(rec.amount, Decimal):
            issues.append(Issue(ERROR, "rule_amount_type", "VALOR DE LA FACTURA debe ser numérico.", row))
        if rec.nit is not None and not rec.nit.full.isdigit():
            issues.append(Issue(ERROR, "rule_nit_text", "CÓDIGO DOCUMENTO debe ser texto numérico.", row))
    return issues


# --------------------------------------------------------------------------- #
# Nivel 3: workbook final
# --------------------------------------------------------------------------- #
def _decimal_from_cell(value: Any) -> Decimal | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return Decimal(repr(value)) if isinstance(value, float) else Decimal(value)


def _cell_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date() if value.time() == datetime.min.time() else None
    return None


def verify_workbook(path: Path, records: list[Record], profile: Profile, expected_total: Decimal) -> list[str]:
    """Reabre el archivo guardado y devuelve la lista de problemas (vacía = correcto)."""
    problems: list[str] = []

    def fail(message: str) -> None:
        if len(problems) < 50:
            problems.append(message)

    try:
        output = openpyxl.load_workbook(path)           # sin data_only: las fórmulas serían visibles
        template = openpyxl.load_workbook(profile.template_path)
    except Exception as exc:  # cualquier fallo al abrir invalida el archivo
        return [f"El workbook generado no se pudo abrir: {type(exc).__name__}"]

    try:
        if output.sheetnames != template.sheetnames:
            fail(f"Las hojas no coinciden con la plantilla: {output.sheetnames} != {template.sheetnames}")
            return problems
        if profile.sheet not in output.sheetnames:
            fail(f"No existe la hoja principal '{profile.sheet}'.")
            return problems
        for name in profile.auxiliary_sheets:
            if name not in output.sheetnames:
                fail(f"Falta la hoja auxiliar '{name}'.")

        _verify_auxiliary_sheets(output, template, profile, fail)

        ws = output[profile.sheet]
        template_ws = template[profile.sheet]
        _verify_structure(ws, template_ws, profile, fail)
        _verify_rows(ws, template_ws, records, profile, expected_total, fail)
    finally:
        output.close()
        template.close()
    return problems


def _verify_auxiliary_sheets(output: Any, template: Any, profile: Profile, fail: Any) -> None:
    for name in template.sheetnames:
        if name == profile.sheet:
            continue
        out_ws, tpl_ws = output[name], template[name]
        if out_ws.sheet_state != tpl_ws.sheet_state:
            fail(f"La hoja '{name}' cambió de visibilidad.")
        out_rows = list(out_ws.iter_rows(values_only=True))
        tpl_rows = list(tpl_ws.iter_rows(values_only=True))
        if out_rows != tpl_rows:
            fail(f"El contenido de la hoja auxiliar '{name}' difiere de la plantilla.")


def _verify_structure(ws: Any, template_ws: Any, profile: Profile, fail: Any) -> None:
    for col in profile.columns:
        actual = ws.cell(profile.header_row, col.index).value
        if actual != col.header:
            fail(f"Encabezado incorrecto en {col.letter}{profile.header_row}: {actual!r} != {col.header!r}.")
    for letter, dim in template_ws.column_dimensions.items():
        out_dim = ws.column_dimensions.get(letter)
        if out_dim is None or bool(out_dim.hidden) != bool(dim.hidden):
            fail(f"La visibilidad/ancho de la columna {letter} cambió respecto a la plantilla.")


def _verify_rows(ws: Any, template_ws: Any, records: list[Record], profile: Profile,
                 expected_total: Decimal, fail: Any) -> None:
    first = profile.first_data_row
    proto_row = profile.style_prototype_row
    fmt_issue = template_ws.cell(proto_row, profile.column("issue_date").index).number_format
    fmt_due = template_ws.cell(proto_row, profile.column("due_date").index).number_format
    if not is_date_format(fmt_issue) or not is_date_format(fmt_due):
        fail("La plantilla no define formato de fecha en FECHA EMISIÓN/FECHA VENCIMIENTO.")

    offset = int(profile.rules["due_offset_days"])
    cols = {c.key: c.index for c in profile.columns}
    last_expected = first + len(records) - 1
    total = Decimal(0)

    for i, rec in enumerate(records):
        r = first + i
        get = lambda key: ws.cell(r, cols[key])  # noqa: E731

        _expect_text(get("document_type"), profile.constants["document_type"], "TIPO DOCUMENTO", r, fail)
        _expect_text(get("country"), profile.constants["country"], "PAÍS", r, fail)
        _expect_text(get("nit"), rec.nit.full if rec.nit else None, "CÓDIGO DOCUMENTO", r, fail)
        _expect_text(get("name"), rec.name, "NOMBRE / RAZÓN SOCIAL", r, fail)
        _expect_text(get("invoice"), rec.invoice, "NÚMERO FACTURA", r, fail)

        issue_cell, due_cell = get("issue_date"), get("due_date")
        issue, due = _cell_date(issue_cell.value), _cell_date(due_cell.value)
        if issue is None or issue_cell.data_type != "d":
            fail(f"Fila {r}: FECHA EMISIÓN no es una fecha real de Excel.")
        elif issue != rec.issue_date:
            fail(f"Fila {r}: FECHA EMISIÓN incorrecta ({issue} != {rec.issue_date}).")
        if due is None or due_cell.data_type != "d":
            fail(f"Fila {r}: FECHA VENCIMIENTO no es una fecha real de Excel.")
        elif issue is not None and due - issue != timedelta(days=offset):
            fail(f"Fila {r}: FECHA VENCIMIENTO no es emisión + {offset} días.")
        if issue_cell.number_format != fmt_issue or not is_date_format(issue_cell.number_format):
            fail(f"Fila {r}: formato de FECHA EMISIÓN distinto al de la plantilla.")
        if due_cell.number_format != fmt_due or not is_date_format(due_cell.number_format):
            fail(f"Fila {r}: formato de FECHA VENCIMIENTO distinto al de la plantilla.")

        term_cell = get("term")
        if isinstance(term_cell.value, bool) or not isinstance(term_cell.value, int) \
                or term_cell.value != profile.constants["term_days"] or term_cell.data_type != "n":
            fail(f"Fila {r}: PLAZO DE VENTA debe ser el entero {profile.constants['term_days']}.")

        amount_cell = get("amount")
        amount = _decimal_from_cell(amount_cell.value)
        if amount is None or amount_cell.data_type != "n":
            fail(f"Fila {r}: VALOR DE LA FACTURA no es numérico.")
        else:
            total += amount
            if rec.amount is not None and amount != Decimal(repr(float(rec.amount))):
                fail(f"Fila {r}: VALOR DE LA FACTURA incorrecto ({amount} != {rec.amount}).")

        for key in ("solunion_code", "debtor_reference", "payment_condition"):
            expected = getattr(rec, key)
            _expect_optional(get(key), expected, profile.column(key).public_name, r, fail)

    if abs(total - expected_total) >= Decimal("0.005"):
        fail(f"La suma de VALOR DE LA FACTURA ({total}) no coincide con el total esperado ({expected_total}).")

    # No deben quedar datos antiguos: nada con valor fuera de las filas escritas.
    for r in range(first, ws.max_row + 1):
        if r <= last_expected:
            continue
        for c in range(1, ws.max_column + 1):
            if ws.cell(r, c).value is not None:
                fail(f"Quedó un valor residual en {ws.cell(r, c).coordinate}.")
                break
    extra_columns = [c for c in range(1, profile.style_columns + 1) if c not in cols.values()]
    for r in range(first, last_expected + 1):
        for c in extra_columns:
            if ws.cell(r, c).value is not None:
                fail(f"Valor inesperado en {ws.cell(r, c).coordinate} (columna fuera de la plantilla).")


def _expect_text(cell: Any, expected: Any, label: str, row: int, fail: Any) -> None:
    value = cell.value
    if not isinstance(value, str) or cell.data_type != "s":
        fail(f"Fila {row}: {label} debe ser texto (tipo actual: {type(value).__name__}/{cell.data_type}).")
    elif expected is not None and value != expected:
        fail(f"Fila {row}: {label} incorrecto ({value!r} != {expected!r}).")


def _expect_optional(cell: Any, expected: str | None, label: str, row: int, fail: Any) -> None:
    if expected is None:
        if cell.value is not None:
            fail(f"Fila {row}: {label} debe quedar vacío.")
    else:
        _expect_text(cell, expected, label, row, fail)

