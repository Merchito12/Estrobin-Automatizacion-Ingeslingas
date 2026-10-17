"""Genera backend/templates/PlantillaDeclaracionVentas_BASE.xlsx a partir de un
archivo ACEPTADO por la plataforma (por ejemplo, el de julio).

Conserva hojas, orden, encabezados, estilos, formatos, anchos, columnas ocultas y
hojas auxiliares. Elimina únicamente los registros de ventas y uniforma el estilo
de la zona de datos al de la primera fila de datos (fila prototipo).

Uso:
    python scripts/build_base_template.py "PlantillaDeclaracionVentas JULIO 2026.xlsx"
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

import openpyxl

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from app.profile import get_profile  # noqa: E402


def build(source: Path, target: Path) -> dict[str, int]:
    profile = get_profile()
    workbook = openpyxl.load_workbook(source, keep_links=False)
    worksheet = workbook[profile.sheet]

    for column in profile.columns:
        actual = worksheet.cell(profile.header_row, column.index).value
        if actual != column.header:
            raise SystemExit(
                f"Encabezado {column.letter}{profile.header_row} inesperado: {actual!r} (esperado {column.header!r})"
            )

    prototype = {
        col: copy(worksheet.cell(profile.style_prototype_row, col)._style)
        for col in range(1, profile.style_columns + 1)
    }
    prototype_height = worksheet.row_dimensions[profile.style_prototype_row].height

    cleared = 0
    for row in range(profile.first_data_row, worksheet.max_row + 1):
        for col in range(1, profile.style_columns + 1):
            cell = worksheet.cell(row, col)
            if cell.value is not None:
                cleared += 1
            cell.value = None
            cell._style = copy(prototype[col])
        worksheet.row_dimensions[row].height = prototype_height

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return {"cleared_cells": cleared, "rows": worksheet.max_row}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    destination = get_profile().template_path
    summary = build(Path(sys.argv[1]), destination)
    print(f"Plantilla base escrita en {destination} ({summary['cleared_cells']} celdas con datos eliminadas)")
